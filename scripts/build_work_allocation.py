"""Create the revised collaborator handoff from saved results and explicit tasks."""
from pathlib import Path
import csv
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]


def main():
    final={r['model']:r for r in csv.DictReader((ROOT/'results/final_evaluation/results.csv').open())}
    audits=list(csv.DictReader((ROOT/'results/explanations/financial_summary.csv').open()))
    current=next(r for r in audits if r['seed']=='11' and r['lag']=='0')
    pages=[]
    def page(title):
        page_={'title':title,'blocks':[]};pages.append(page_);return page_['blocks']
    def para(blocks,text,bold=False):blocks.append({'type':'p','text':text,'bold':bold})
    def heading(blocks,text):blocks.append({'type':'h','text':text})
    def task(blocks,title,work,deliverable,acceptance):
        heading(blocks,title);para(blocks,work);para(blocks,'Deliverable: '+deliverable);para(blocks,'Complete when: '+acceptance)
    p=page('RiskTrace research progress and revised work allocation')
    para(p,'For Ayush Sinha and Atharva Sheersh Pandey | 27 October 2026')
    para(p,'Ayush will now own finance verification and a defined ML workstream. Atharva will retain the primary pipeline and research integration. This replaces the earlier allocation that assigned no ML work to Ayush. The immediate goal is to verify and strengthen an exploratory research study before deciding its publication direction.')
    heading(p,'Work completed on Atharva side')
    para(p,'The repository contains the official-source collection, allocation of 807 banks to each person, validated canonical panel, chronological splits, training-only preprocessing, logistic regression, XGBoost, GRU, static GAT, temporal GAT plus GRU, graph controls, explanation audits and robustness analysis. There are 297 development fits and 23 final fits. These are repository artifacts produced with AI assistance; they do not establish unrecorded human review or authorship contributions.')
    para(p,'The dataset has 51,339 eligible bank-quarter reports from 1,614 banks during 2010Q1 to 2025Q4. Annual rolling validation uses 2012 to 2020 and only six distinct failed banks. The observable test contains 12,783 reports and four distinct failed banks. The source is a centrally assembled official panel; Ayush\'s separate submission has not been received or verified.')
    heading(p,'What the actual results mean')
    para(p,f"At the unchanged validation-derived thresholds, XGBoost and every graph variant detect zero of four failed test banks. Logistic regression detects one, with achieved test false-alarm rate {float(final['logistic']['false_alarm_rate'])*100:.3f}%. The graph models have not shown useful warning superiority.")
    para(p,f"Current financial-input edits change the top-ranked group in {float(current['top1_disagreement'])*100:.2f}% of complete audits, but the mean probability gap is only {float(current['mean_gap'])*100:.6f} percentage points. All edge masks explain nonfailure at the 0.5 class boundary. In 82 of 500 bootstrap draws AP is undefined. Financial usefulness, high novelty and Q1 readiness are unestablished.")
    para(p,'The manuscript, claim map, saved checkpoints, case packets and reproduction tools are assembled. Fresh-environment inference has been exercised; the final verification report accompanies the handoff. Author review, finance validation, the final venue and submission approval are pending.')

    p=page('Supported novelty and research questions')
    para(p,'The defensible contribution is a dependency-aware empirical audit of bank-forecast explanations. It compares the same financial edit with the original graph retained and with feature-derived links and affected temporal states recomputed, while keeping fitted weights and preprocessing unchanged. These answer two legitimate explanation questions. A difference does not prove that either explanation is wrong or causal.')
    p.append({'type':'table','headers':['Research question','What this study answers','Gap addressed and limit'],'rows':[
        ['Do financial dependencies change explanation conclusions?','Yes for some ranking conclusions: current top-one disagreement is 13.49% among complete audits; the mean probability gap is small.','Makes raw inputs, derived ratios, graph links and temporal replay explicit. Does not establish a first-ever method.'],
        ['Do similarity graphs improve failure warnings?','No useful gain is demonstrated under this bounded protocol; all graph variants miss the four test failures at fixed thresholds.','Adds transparent tabular, no-graph and graph controls. Four events cannot settle graph usefulness generally.'],
        ['Do rankings imply materially different alerts?','The two audit targets cross the same alert threshold differently in about 0.88% of donor-matched current-edit cases.','Separates ranking changes from probability and alert changes. Economic materiality and causal validity remain unproved.'],
        ['How reliable is sparse-event evidence?','Six development and four test failures limit selection and uncertainty; 82 AP resamples are undefined.','Records label maturity, event counts, invalid probes and dependence-aware sensitivity. This is methodological discipline, not inherently novel.']
    ]})
    para(p,'An additional finding is a target mismatch: a model can generate low-probability operational alerts while a binary-class explainer still explains nonfailure. This was observed here; a general solution has not been developed or validated.')
    para(p,'The study does not fill the gaps of causal contagion identification, measured interbank exposure reconstruction, strict historical-vintage backtesting or a deployable early-warning system. It also does not establish that GAT plus GRU, attention plots or graph rebuilding are new.')
    para(p,'Closest overlaps to check: CF3 financial factual/counterfactual explanation; GInX-Eval, DSE and Robust Fidelity evaluation; temporal bank graph forecasting. The comparison matrix exists, but close-method reproduction and proof of a distinct methodological advance remain pending. Sources: docs/literature_matrix.csv, docs/final_novelty_and_research_questions.md and paper/references.bib.')

    p=page('Ayush tasks for finance and baseline modeling')
    task(p,'1 Verify the allocated banks and financial definitions',
         'Use data/bank_allocation.csv to reconcile your 807 banks with source IDs and report keys. If you have a separate collection, provide its raw files, source URLs, retrieval dates, units and missing-value flags. Review EQTOT versus EQ, deposits scope, failure versus assistance transactions, mergers and renamed entities. Record disagreements without editing the frozen panel.',
         'A source and entity discrepancy ledger, plus your actual collection manifest or an explicit statement that no separate collection is supplied.',
         'Each issue cites a CERT, report date, source and proposed resolution; Atharva can reproduce it.')
    task(p,'2 Reproduce logistic regression and XGBoost independently',
         'Own the ML implementation of a separate baseline replication runner. Use the approved annual 2012 to 2020 folds, original training keys and supplied training-fitted preprocessing. Take the already selected hyperparameters from selected_models.json. Start with the 2012 fold and one seed, then run all nine folds. Use seed11 for logistic and seeds11,29,47 for XGBoost. Do not select parameters from the final test or overwrite primary results.',
         'Your runner, environment details, fold predictions and a comparison report under results/ayush_replication/. Include AP, Brier, validation-selected thresholds, per-quarter recall and distinct-event warnings.',
         'Differences from saved development predictions/metrics are reconciled within a stated numerical tolerance, with a reason for every sample or implementation difference.')
    para(p,'Study src/risktrace/predictive.py, analysis.py and metrics.py before implementing. Preserve horizon maturity at each forecast origin. Positive quarters are repeated windows, not additional independent failures. Atharva reviews the runner before the full replication. This is new assigned work; no Ayush implementation or completion is currently claimed.')
    para(p,'Working estimate: source review and a one-fold ML smoke run in the first three working days; full fixed-setting replication by the end of the first working week, subject to setup and discrepancies. These are planning targets, not recorded commitments.')

    p=page('Ayush tasks for explanation analysis and interpretation')
    task(p,'3 Own development calibration and explanation diagnostics',
         'Using your independent out-of-fold predictions, implement reliability tables, Brier calculation, false-alarm accounting and event-level warning reconstruction. Add an exploratory tabular explanation comparison: logistic coefficient contributions and XGBoost TreeSHAP on the same development records. Confirm the additive prediction scale, transformed-input units and comparable record selection. Explain what each method represents and avoid treating attribution as a causal effect.',
         'A notebook or script, development-only diagnostic tables and five to ten correctly sourced example-bank explanations. Use an outcome-independent sample for aggregate summaries.',
         'Predictions and contribution totals reconcile to each model on the documented scale; metrics match the independent evaluation; no claims are selected from attractive examples.')
    task(p,'4 Design a development experiment for explaining alerts',
         'Study why the original binary-class masks explain nonfailure while operational alert thresholds are much lower. With Atharva, specify whether the new target is failure probability or distance above the frozen alert threshold. Prototype the comparison only on development data, retaining equal explanation budgets, random controls and valid financial probes. Label it exploratory and keep its outputs separate. Do not claim this solves the mismatch until tested.',
         'A short ML experiment specification and a development prototype with unit checks, response curves, random controls and failure cases.',
         'The target, output scale, thresholds, sampling, budgets and temporal replay are explicit. Any new confirmatory claim requires a newly agreed evaluation protocol; the inspected test cannot become a fresh holdout.')
    task(p,'5 Review cases and sharpen the literature comparison',
         'Review the four proposed failure/control packets, distinguish Republic Bank from First Republic Bank, and check control comparability. Cite primary FDIC filings/releases for financial interpretation. Compare the exact research question with CF3, GInX-Eval, DSE and Robust Fidelity; record what each actually evaluates.',
         'A sourced case commentary and close-paper comparison with overlap, differences and unresolved questions.',
         'Retrospective narrative is separate from prediction evidence; unsupported claims are removed. Approval or submission is not inferred from completion of a task.')

    p=page('Atharva tasks and chronological handoff')
    para(p,'Atharva retains primary data integration, frozen models, graph and temporal code, dependency replay, reproducibility tooling and the manuscript. Ayush now shares ML implementation through the separate workstream above; this is not a transfer of the original primary-result provenance.')
    p.append({'type':'table','headers':['Order','Owner','Next action and required output'],'rows':[
        ['1','Atharva','Hand over the fixed dataset manifest, split keys, selected settings, primary results and case packets. Preserve the original version.'],
        ['2','Ayush','Reconcile allocated data and source/entity definitions; build one-fold independent baseline replication.'],
        ['3','Atharva','Review the smoke run, resolve implementation discrepancies and document any proposed data correction as a new version.'],
        ['4','Ayush','Complete fixed-setting baseline reproduction and development calibration/attribution diagnostics.'],
        ['5','Both','Specify the alert-explanation development experiment; Ayush prototypes it and Atharva verifies temporal/dependency handling.'],
        ['6','Ayush','Provide primary-source case interpretation and close-method overlap review.'],
        ['7','Atharva','Integrate justified findings, rerun affected checks and update the manuscript with exact result provenance. Preserve all negative results.'],
        ['8','Both','Decide whether the evidence supports a scoped audit/negative-result paper or requires a redesigned study with new evaluation data.'],
        ['9','Both','Verify current venue requirements and ranking evidence, agree authorship and disclosures, review the exact version and approve before any submission.']
    ]})
    heading(p,'Repository and review rules')
    para(p,'Keep Ayush replication and exploratory outputs in clearly separate folders. Use detailed commits with purpose, changes, validation and limitations; review before merging. Never replace frozen predictions, thresholds or the analysis plan to improve a test result. A genuine correction needs a logged reason, version and downstream rerun.')
    para(p,'For every handoff provide the commit SHA, commands, environment, dataset hash, output paths, checks passed and unresolved issues. No author order, affiliation, finance approval, Q1 status, submission or acceptance is assumed. The approximate 16-week target must be revised around the current evidence; publication timing is controlled by the venue.')
    para(p,'Repository: https://github.com/atharvasheersh/RiskTrace | Primary evidence: results/final_evaluation/, results/explanations/, results/robustness/, results/case_studies/ and docs/claim_evidence_map.csv.')

    output=ROOT/'output/documents';output.mkdir(parents=True,exist_ok=True)
    (output/'work_allocation_content.json').write_text(json.dumps(pages,indent=2)+'\n',encoding='utf8')
    markdown=[]
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
    for name,size in [('Title',23),('Heading 1',16),('Heading 2',12)]:
        style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.keep_with_next=True;style.paragraph_format.space_after=Pt(9)
        style.paragraph_format.space_before=Pt(8)
    doc.styles['Title'].paragraph_format.space_before=Pt(0)
    for i,page_ in enumerate(pages):
        if i:doc.add_page_break()
        doc.add_paragraph(page_['title'],'Title' if i==0 else 'Heading 1');markdown+=['# '+page_['title'],'']
        for block in page_['blocks']:
            kind=block['type']
            if kind in ['p','h']:
                par=doc.add_paragraph(block['text'],'Heading 2' if kind=='h' else 'Normal')
                if block.get('bold'):
                    for run in par.runs:run.bold=True
                markdown += [('## ' if kind=='h' else '')+block['text'],'']
            else:
                table=doc.add_table(rows=1,cols=len(block['headers']));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
                widths=[2.0,2.25,2.75] if i==1 else [.5,.85,5.65]
                for j,text in enumerate(block['headers']):table.rows[0].cells[j].text=text
                for row in block['rows']:
                    cells=table.add_row().cells
                    for j,text in enumerate(row):cells[j].text=text
                for ri,row in enumerate(table.rows):
                    if ri==0:
                        repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
                    no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
                    for j,cell in enumerate(row.cells):
                        cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243F58' if ri==0 else ('F1F4F6' if ri%2 else 'FFFFFF'));pr.append(shade)
                        borders=OxmlElement('w:tcBorders')
                        for side in ['top','left','bottom','right']:
                            edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                        pr.append(borders);margins=OxmlElement('w:tcMar')
                        for side in ['top','left','bottom','right']:
                            margin=OxmlElement('w:'+side);margin.set(qn('w:w'),'85');margin.set(qn('w:type'),'dxa');margins.append(margin)
                        pr.append(margins)
                        for par in cell.paragraphs:
                            par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.03
                            for run in par.runs:run.font.size=Pt(10);run.bold=ri==0;run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
                doc.add_paragraph().paragraph_format.space_after=Pt(2)
                markdown+=['| '+' | '.join(block['headers'])+' |','| '+' | '.join(['---']*len(block['headers']))+' |']
                markdown+=['| '+' | '.join(row)+' |' for row in block['rows']];markdown+=['']
    doc.core_properties.title=pages[0]['title'];doc.core_properties.subject='Revised allocation with Ayush ML tasks and research status'
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    # Word's bundled Title style can carry a blue bottom rule. Remove that
    # inherited border as well as any direct paragraph border before saving.
    for root in [doc.element,doc.styles.element]:
        for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
    path=output/'RiskTrace_Revised_Work_Allocation_Ayush_Atharva.docx';doc.save(path)
    (ROOT/'docs/revised_work_allocation.md').write_text('\n'.join(markdown),encoding='utf8',newline='\n')
    print(path)


if __name__=='__main__':main()
